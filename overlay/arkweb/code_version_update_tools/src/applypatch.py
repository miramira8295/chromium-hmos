#!/usr/bin/env python3
# -*- coding: utf-8 -*-

# Copyright (c) 2025 Huawei Device Co., Ltd.
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

import os
import sys
import pickle
import datetime

import util
import config
import subprocess

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

import shutil


def get_folder_differ_map():
    dump_filename = os.path.join(config.root, 'result', 'folder_differ_map.pkl')
    with open(dump_filename, "rb") as f:
        folder_differ_map = pickle.load(f)
        return folder_differ_map
    return {}


def dump_patched_map(folder_differ_map, suffix):
    dump_filename = os.path.join(config.root, 'result', f'folder_patched_map_{suffix}.pkl')
    with open(dump_filename, "wb") as f:
        pickle.dump(folder_differ_map, f)


def check_line_ending_with_file_command(file_path):
    # 执行 file 命令并捕获输出
    result = subprocess.run(
        ["file", file_path],
        capture_output=True,
        text=True,
        check=True
    )
    output = result.stdout.lower()
    if "crlf" in output:
        return "DOS"
    else:
        return "Unix"


def apply_patch(folder_differ_map, work_path):
    for (dir_name, file_differ_list) in folder_differ_map.items():
        for file_differ_obj in file_differ_list:
            if file_differ_obj.is_ignored_file:
                print(f"skip: {file_differ_obj.save_filename}")
                continue
            result_diff = file_differ_obj.save_filename

            if file_differ_obj.status == "[D]":
                continue
            elif file_differ_obj.status == "[A]":
                cmd = f"cp -r {file_differ_obj.right_file} {os.path.join(work_path, file_differ_obj.project_name, file_differ_obj.relative_filename)}"
            elif file_differ_obj.status == "[M]":
                target_file = os.path.join(work_path, file_differ_obj.project_name, file_differ_obj.relative_filename)
                cmd = f"patch --ignore-whitespace -F 2 --force --binary {target_file} < {result_diff}"
            stderr, result = util.exec_and_stderr_stdout(cmd, True)

            count_total_hunk = 0
            count_fail_hunk = 0
            rej_filename = os.path.join(work_path, file_differ_obj.project_name, f"{file_differ_obj.relative_filename}.rej")
            try:
                with open(result_diff, 'r', encoding='utf-8') as f:
                    tmp_data = f.readlines()
                    for i in tmp_data:
                        if '@@' in i:
                            count_total_hunk += 1
            except OSError:
                pass
            except UnicodeDecodeError:
                with open(result_diff, 'r', encoding='GB2312') as f:
                    tmp_data = f.readlines()
                    for i in tmp_data:
                        if '@@' in i:
                            count_total_hunk += 1
            if os.path.exists(rej_filename):
                with open(rej_filename, 'r') as f:
                    tmp_data = f.readlines()
                    for i in tmp_data:
                        if '@@' in i:
                            count_fail_hunk += 1
            file_differ_obj.count_total_hunk = count_total_hunk
            file_differ_obj.count_fail_hunk = count_fail_hunk

            if stderr != "" or result.find("FAILED") != -1 or result.find("**") != -1:
                output = f"\033[5;31;40m[PATCH FAILED] \033[0m[{file_differ_obj.status}]{stderr}\n{result}"
                print(output, end="")
                print(cmd)
                record_info[1] = record_info[1] + 1
                rej_filename = os.path.join(work_path, file_differ_obj.project_name, f"{file_differ_obj.relative_filename}.rej")
                if not os.path.exists(rej_filename):
                    rej_filename = result_diff
                    try:
                        with open(rej_filename, 'r') as f:
                            tmp_data = f.readlines()
                            for i in tmp_data:
                                if '@@' in i:
                                    count_fail_hunk += 1
                        file_differ_obj.count_fail_hunk = count_fail_hunk
                    except OSError:
                        pass
                command = f"grep -E '^[+-]' {rej_filename} | wc -l "
                file_differ_obj.merge_faild_line = int(util.exec_and_output(command, False)) - 2 # 减掉.rej文件开头的+++ ---两行
                continue
            file_differ_obj.set_merge_status(True)
            record_info[0] = record_info[0] + 1
            util.INFO(f"[PATCH PASSED] [{file_differ_obj.status}]{result_diff}")


def export_all_diff(folder_differ_map, suffix):
    dump_patched_map(folder_differ_map, suffix)

    # source_right_for_apply的目录不需要生成excel表格
    if suffix == 'source_right_for_apply':
        return

    # 创建Workbook
    wb = Workbook()
    ws = wb.active
    # 写入标题
    headers = ["文件名", "总行数", "失败行数", "总Hunk数", "失败Hunk数", "是否成功应用", "类型"]
    ws.append(headers)

    # 设置标题样式
    for col in range(1, len(headers)+1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    
    for (dir_name, file_differ_list) in folder_differ_map.items():
        for file_differ_obj in file_differ_list:
            filename = os.path.join("src", file_differ_obj.project_name, file_differ_obj.relative_filename)
            filename = filename.replace("//", "/")
            merge_status = file_differ_obj.merge_status
            status = file_differ_obj.status
            if file_differ_obj.is_ignored_file:
                status = "[IGNORED]ignored"
            ws.append([os.path.normpath(filename), int(file_differ_obj.diff_line), int(file_differ_obj.merge_faild_line), int(file_differ_obj.count_total_hunk), int(file_differ_obj.count_fail_hunk), merge_status, status])
    current_time = datetime.datetime.now()
    formatted_time = current_time.strftime("%Y-%m-%d-%H-%M-%S")
    wb.save(f"result/report_{formatted_time}.xlsx")
    shutil.copy(f"result/report_{formatted_time}.xlsx", f"result/report.xlsx")


def main(work_path, branch):
    util.INFO("===========apply patch running=============")
    util.INFO(f"work dir: {work_path}")

    current_time = datetime.datetime.now()
    formatted_time = current_time.strftime("%Y-%m-%d-%H-%M-%S")

    cmd = f"cd {work_path}"
    stderr, result = util.exec_and_stderr_stdout(cmd, True)
    cmd = f"cd {work_path} && git reset --hard"
    stderr, result = util.exec_and_stderr_stdout(cmd, True)
    cmd = f"cd {work_path} && git clean -fxd"
    stderr, result = util.exec_and_stderr_stdout(cmd, True)
    cmd = f"cd {work_path} && git checkout {branch}"
    stderr, result = util.exec_and_stderr_stdout(cmd, True)
    cmd = f"cd {work_path} && git branch -D {branch}_for_process"
    stderr, result = util.exec_and_stderr_stdout(cmd, True)
    cmd = f"cd {work_path} && git checkout -b {branch}_for_process"
    stderr, result = util.exec_and_stderr_stdout(cmd, True)
    cmd = f"cd {work_path} && git branch -u origin/{branch}"
    stderr, result = util.exec_and_stderr_stdout(cmd, True)
    util.ERROR(stderr)

    diff_map = get_folder_differ_map()

    apply_patch(diff_map, work_path)
    if 'source_right_for_apply' in work_path.split(os.sep):
        export_all_diff(diff_map, 'source_right_for_apply')
    elif 'target' in work_path.split(os.sep):
        export_all_diff(diff_map, 'target')
    else:
        print('apply.py 未知错误，系统退出', file=sys.stderr)
        sys.exit(-1)
    util.INFO("-----------------------------------")
    util.INFO("")
    util.INFO(f"[PASSED]: {record_info[0]}")
    util.ERROR(f"[FAILED]: {record_info[1]}")


if __name__ == '__main__':
    # source_right
    type_source_right, branch_source_right, url_source_right = config.remote_source_right_path
    work_path_source_right = os.path.join(config.source_right_for_apply_path, 'src')
    record_info = [0, 0]
    main(work_path_source_right, branch_source_right)

    # target
    type_target, branch_target, url_target = config.remote_target_path
    work_path_target = os.path.join(config.target_path, 'chromium', 'src')
    record_info = [0, 0]
    main(work_path_target, branch_target)
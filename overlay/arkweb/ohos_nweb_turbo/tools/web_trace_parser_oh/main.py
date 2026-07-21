# This is a sample Python script.

# Press Shift+F10 to execute it or replace it with your code.
# Press Double Shift to search everywhere for classes, files, tool windows, actions, and settings.

import os
from openpyxl import Workbook
from src.tools.excel_writer import create_sheet, write_summary_excel
from src.tools.file_handler import read_trace, create_archive_folder

from src.network.index import (
    parse_network_info,
    get_network_calculated_info,
    get_network_statistic_info_map,
    merge_network_timing_interval
)
from src.process.index import (
    render_process_sleeping_info_write_to_excel
)
from src.main.index import (
    get_basic_trace_info,
    parse_and_statistic_main_process,
    trace_statistic_info_write_to_excel
)
from src.main.interface.main_process_trace_info_statistic import (
    MainProcessTraceStatisticMap,
    MainProcessTraceStatistic,
    FirstFrameInfoMap,
    FirstFrameInfo
)
from src.network.interface.network_timing_info import NetWorkStatisticMap
from src.process.interface.render_process_statistic import (
    RenderProcessSleepingStatisticMap,
    RenderProcessSleepingStatistic
)
from src.main.interface.statistic_unit import StatisticUnit

# 配置项
start_step = ["DispatchTouchEvent", "type=1"]
end_step = "largestContentfulPaint"
is_end_by_white_screen = False
filter_url = [
    "https://browsercfg-drcn.cloud.dbankcloud.cn/security/v1/oh/securityurls",
    "https://newsfeed-drcn.cloud.dbankcloud.cn/collector/v1/landing_page_event"
]

# trace输入目录
INPUT_DIR = 'input'
# excel输出目录
OUTPUT_DIR = 'output'

# excel sheet name
network_info_sheet = "子资源数据"
network_calculated_info_sheet = "子资源耗时统计"
render_process_sleeping_info_sheet = "render进程空闲相关信息"
main_process_summary_sheet = "主流程拆解统计"


def parse_nweb_trace_time(basic_trace_info, lines: list[str], excel_path: str):
    basic_trace_info.show_info()
    # 监测trace关键点，关键节点不全无法正常解析
    # if basic_trace_info.need_break():
    #     return

    # 只统计到首帧结束
    if is_end_by_white_screen:
        if basic_trace_info.first_paint != 0:
            basic_trace_info.end_time = basic_trace_info.first_paint
        elif basic_trace_info.first_swap_buffer_after_commit != 0:
            basic_trace_info.end_time = basic_trace_info.first_swap_buffer_after_commit
        else:
            print("未找到fp 且 未找到主资源下载后的首帧。")

    # basic_trace_info.end_time = basic_trace_info.last_swap_buffer_time
    workbook = Workbook()

    workbook.remove(workbook['Sheet'])
    create_sheet(workbook, excel_path, main_process_summary_sheet)
    create_sheet(workbook, excel_path, network_info_sheet)
    create_sheet(workbook, excel_path, network_calculated_info_sheet)
    create_sheet(workbook, excel_path, render_process_sleeping_info_sheet)

    # 解析主流程
    main_process_init_info, main_process_trace_statistic = parse_and_statistic_main_process(lines, basic_trace_info)

    # 解析网络子资源原始数据
    network_info_map = parse_network_info(
        excel_path,
        network_info_sheet,
        lines,
        filter_url,
        basic_trace_info
    )

    merge_network_timing_interval(
        excel_path,
        network_info_sheet,
        network_info_map
    )

    # 计算网络子资源分段耗时
    network_calculated_info_map = get_network_calculated_info(
        excel_path,
        network_calculated_info_sheet,
        network_info_map
    )

    # 解析子资源统计数据
    network_statistic = get_network_statistic_info_map(
        network_calculated_info_map,
    )

    # 解析空闲时间
    render_process_sleeping_statistic = render_process_sleeping_info_write_to_excel(
        excel_path,
        render_process_sleeping_info_sheet,
        lines,
        basic_trace_info.render_process_id,
        basic_trace_info.start_time,
        basic_trace_info.end_time,
        # sub_resource_info_map
    )

    # 计算主流程中的各项耗时占比
    percentage_statistic(
        main_process_trace_statistic,
        render_process_sleeping_statistic
    )

    # # 首帧分段耗时
    # first_frame_timing = FirstFrameInfo(
    #     main_process_trace_statistic.touch_up_to_first_frame_time,
    #     main_process_trace_statistic.first_frame_evaluate_script_time,
    #     main_process_trace_statistic.first_frame_v8_call_function_time,
    #     network_statistic.end_before_first_frame_num
    # )

    # 合并多模块统计数据 总结sheet内容
    statistic_data_map_list = [
        StatisticUnit('主流程', MainProcessTraceStatisticMap, main_process_trace_statistic),
        StatisticUnit('子资源', NetWorkStatisticMap, network_statistic),
        StatisticUnit('render进程空闲', RenderProcessSleepingStatisticMap, render_process_sleeping_statistic),
        # StatisticUnit('首帧前', FirstFrameInfoMap, first_frame_timing)
    ]

    # 将多模块统计数据写入excel第一个sheet中
    trace_statistic_info_write_to_excel(
        excel_path,
        main_process_summary_sheet,
        statistic_data_map_list
    )

    return statistic_data_map_list


def percentage_statistic(
        main_process_trace_statistic: MainProcessTraceStatistic,
        render_process_sleeping_statistic: RenderProcessSleepingStatistic
):
    if main_process_trace_statistic.touch_up_to_lcp_time == 0:
        print("trace解析异常， 请确认trace内容中开始和结束点是否完整")
        return

    if render_process_sleeping_statistic.total_sleeping_time == 0:
        print("空闲时间解析异常，无法计算占比， 请确认trace内容中render进程信息是否完整")
        return

    # 解析网络耗时占比 当前先默认空闲时间极大概率是网络造成的
    main_process_trace_statistic.network_percentage = round(
        render_process_sleeping_statistic.total_sleeping_time / main_process_trace_statistic.touch_up_to_lcp_time,
        4
    )
    print(f'网络耗时占比: {main_process_trace_statistic.network_percentage}')

    # 计算js执行占比
    # 当前认为空闲时间主要是网络造成的，出去空闲时间（网络）、必要的渲染绘制耗时、和render进程阻塞时长，剩下的大部分都是js耗时
    # 计算公式：总时间 - render空闲 - 150（大多数页面渲染绘制耗时） - render阻塞耗时长（commit）
    js_time_with_rendering = (
            main_process_trace_statistic.touch_up_to_lcp_time -
            render_process_sleeping_statistic.total_sleeping_time -
            main_process_trace_statistic.commit_time
    )

    if js_time_with_rendering <= 150:
        print('js执行 + 渲染绘制时长 <= 150， js执行时长占比包含渲染绘制时间')
        main_process_trace_statistic.js_percentage = round(
            js_time_with_rendering / main_process_trace_statistic.touch_up_to_lcp_time,
            4
        )
        print(f'js耗时占比: {main_process_trace_statistic.js_percentage}')

    else:
        main_process_trace_statistic.js_percentage = round(
            (js_time_with_rendering - 150) / main_process_trace_statistic.touch_up_to_lcp_time,
            4
        )
        print(f'js耗时占比: {main_process_trace_statistic.js_percentage}')


def main():
    if not os.path.exists(INPUT_DIR):
        print(f'不存在{INPUT_DIR}目录，请先新建{INPUT_DIR}目录，并放入需要解析的trace')
        return

    file_list = os.listdir(INPUT_DIR)
    summary_list: list[list[StatisticUnit]] = []
    summary_key_map = {**MainProcessTraceStatisticMap, **NetWorkStatisticMap, **RenderProcessSleepingStatisticMap}
    data_map_list = []
    for file_name in file_list:
        if not any(suffix in file_name for suffix in ['.txt', '.ftrace']):
            continue
        lines = read_trace(INPUT_DIR, file_name)
        if not lines or len(lines) == 0:
            print(f'trace: {file_name} 内容为空')
            continue

        excel_path = create_archive_folder(file_name, OUTPUT_DIR)
        basic_trace_info = get_basic_trace_info(lines, start_step, end_step, file_name)
        statistic_data_map_list = parse_nweb_trace_time(basic_trace_info, lines, excel_path)
        summary_list.append(statistic_data_map_list)
        data_map = {}
        for unit in statistic_data_map_list:
            data_map.update(unit.data_map.__dict__)
        data_map_list.append(data_map)
        print(f"解析 {file_name} 完成")

    print('生成汇总表...')
    write_summary_excel(os.path.join(OUTPUT_DIR, '汇总.xlsx'), '汇总', summary_key_map, data_map_list)

    print("汇总表生成完成")


# Press the green button in the gutter to run the script.
if __name__ == '__main__':
    main()

# See PyCharm help at https://www.jetbrains.com/help/pycharm/
